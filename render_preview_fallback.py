from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


BOOK = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Master_Recon_Summary_Redesigned.xlsx")
OUT = Path("/tmp") / f"recon_rendered_{BOOK.stem}"
OUT.mkdir(parents=True, exist_ok=True)
wb = load_workbook(BOOK, data_only=True)


def rgb(color, default):
    if color and color.type == "rgb" and color.rgb:
        value = color.rgb[-6:]
        try:
            return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))
        except ValueError:
            pass
    return default


def font(size, bold=False):
    regular = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    try:
        return ImageFont.truetype(bold_path if bold else regular, max(7, int(size)))
    except OSError:
        return ImageFont.load_default()


def display_value(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        fmt = cell.number_format or ""
        if "%" in fmt:
            return f"{value:.1%}"
        if "GH¢" in fmt:
            if '0.00,,"M"' in fmt:
                return f"GH¢{value / 1_000_000:,.2f}M"
            return f"-GH¢{abs(value):,.2f}" if value < 0 else f"GH¢{value:,.2f}"
        if "txns" in fmt:
            return f"{value:,.0f} txns"
        if value == int(value):
            return f"{value:,.0f}"
        return f"{value:,.2f}"
    if hasattr(value, "strftime"):
        return value.strftime("%d %b %Y")
    return str(value)


def render_sheet(ws, max_col, max_row, filename, scale=0.72):
    col_widths = []
    for col in range(1, max_col + 1):
        dim = ws.column_dimensions[get_column_letter(col)]
        width = 0 if dim.hidden else (dim.width or 11) * 7.4 * scale
        col_widths.append(width)
    row_heights = []
    for row in range(1, max_row + 1):
        dim = ws.row_dimensions[row]
        height = 0 if dim.hidden else (dim.height or 18) * 1.33 * scale
        row_heights.append(height)
    xs = [0]
    for width in col_widths:
        xs.append(xs[-1] + width)
    ys = [0]
    for height in row_heights:
        ys.append(ys[-1] + height)
    image = Image.new("RGB", (max(1, int(xs[-1])), max(1, int(ys[-1]))), "white")
    draw = ImageDraw.Draw(image)

    merged_lookup = {}
    for merged in ws.merged_cells.ranges:
        if merged.max_col > max_col or merged.max_row > max_row:
            continue
        merged_lookup[(merged.min_row, merged.min_col)] = merged
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                if (row, col) != (merged.min_row, merged.min_col):
                    merged_lookup[(row, col)] = None

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            if (row, col) in merged_lookup and merged_lookup[(row, col)] is None:
                continue
            cell = ws.cell(row, col)
            merged = merged_lookup.get((row, col))
            end_col = merged.max_col if merged else col
            end_row = merged.max_row if merged else row
            x0, y0 = xs[col - 1], ys[row - 1]
            x1, y1 = xs[end_col], ys[end_row]
            fill = rgb(cell.fill.fgColor, (255, 255, 255)) if cell.fill.fill_type else (255, 255, 255)
            draw.rectangle((x0, y0, x1, y1), fill=fill)
            border = (228, 222, 209)
            if any(side is not None and side.style for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom)):
                draw.rectangle((x0, y0, x1, y1), outline=border, width=1)
            text = display_value(cell)
            if not text:
                continue
            color = rgb(cell.font.color, (24, 48, 45))
            fnt = font((cell.font.sz or 10) * scale * 1.28, bool(cell.font.bold))
            pad = 4 * scale
            bbox = draw.textbbox((0, 0), text, font=fnt)
            tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            align = cell.alignment.horizontal or "left"
            if align == "center":
                tx = x0 + (x1 - x0 - tw) / 2
            elif align == "right":
                tx = x1 - tw - pad
            else:
                tx = x0 + pad
            ty = y0 + (y1 - y0 - th) / 2
            if tw > x1 - x0 - 2 * pad:
                text = text[:max(1, int((x1 - x0) / max(5, (cell.font.sz or 10) * scale)))]
            draw.text((tx, ty), text, fill=color, font=fnt)

    if ws.title == "Master Summary" and ws["X2"].value is not None:
        # Mirror the native chart's information in the fallback preview.
        left, top, right, bottom = xs[2], ys[36], xs[11], ys[48]
        draw.rectangle((left, top, right, bottom), fill=(255, 255, 255))
        values = [(ws[f"X{row}"].value, ws[f"Y{row}"].value or 0) for row in range(2, 6)]
        label_font = font(9 * scale, True)
        value_font = font(8 * scale, False)
        max_bar = right - left - 170 * scale
        for idx, (name, value) in enumerate(values):
            y = top + 18 * scale + idx * 33 * scale
            draw.text((left + 10 * scale, y), str(name), fill=(24, 48, 45), font=label_font)
            bx = left + 100 * scale
            draw.rounded_rectangle((bx, y, bx + max_bar, y + 10 * scale), radius=4, fill=(239, 235, 225))
            draw.rounded_rectangle((bx, y, bx + max_bar * value, y + 10 * scale), radius=4, fill=(214, 8, 107))
            draw.text((right - 52 * scale, y - 2 * scale), f"{value:.1%}", fill=(143, 138, 124), font=value_font)

    image.save(OUT / filename)


if "Master Summary" in wb.sheetnames:
    render_sheet(wb["Master Summary"], 21, 86, "master_summary.png", 0.78)
    for sheet_name in wb.sheetnames:
        if "Not Found" not in sheet_name:
            continue
        ws = wb[sheet_name]
        filename = f"{sheet_name.lower().replace(' ', '_')}.png"
        render_sheet(ws, min(12, ws.max_column), min(102, ws.max_row), filename, 0.9)
    if "Source Data" in wb.sheetnames:
        render_sheet(wb["Source Data"], 7, 42, "source_data.png", 1.15)
    if "Checks" in wb.sheetnames:
        render_sheet(wb["Checks"], 8, 17, "checks.png", 1.0)
else:
    render_sheet(wb["Monthly Dashboard"], 21, 52, "monthly_dashboard.png", 0.78)
    render_sheet(wb["Run Status"], 6, 13, "run_status.png", 1.0)
    render_sheet(wb["Collection Summary"], 10, 11, "collection_summary.png", 0.9)
    render_sheet(wb["Disbursement Summary"], 10, 4, "disbursement_summary.png", 0.9)
    render_sheet(wb["Wallet Summary"], 6, 4, "wallet_summary.png", 1.0)
    render_sheet(wb["Granular Detail"], 3, 62, "granular_detail.png", 1.0)
    if "Monthly Checks" in wb.sheetnames:
        render_sheet(wb["Monthly Checks"], 7, 11, "monthly_checks.png", 1.0)
print(OUT)
